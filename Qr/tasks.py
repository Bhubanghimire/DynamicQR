from celery import shared_task
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from subscriptions.models import Subscription

from .models import QRCode, QRImportJob

def get_importer(qr_type):
    qr_type_name = qr_type.id
    print(qr_type_name, str(qr_type_name) == "c5e4c79e-6679-4ef6-ac1c-4b10b7ebbf62")


    if str(qr_type_name) == "221426fb-63e6-4bd3-9d7a-1729d5b53fa5":
        from .importers.wifi import WiFiImporter

        return WiFiImporter()

    if str(qr_type_name) == "80fe2eb5-3a8b-4049-bd21-96bbfd39e6e1":
        from .importers.email import EmailImporter

        return EmailImporter()

    if str(qr_type_name) == "c5e4c79e-6679-4ef6-ac1c-4b10b7ebbf62":
        from .importers.events import EventImporter

        return EventImporter()

    if str(qr_type_name) == "6a16a631-895d-4aa0-b072-ca9cc26a6981":
        from .importers.googlemap import GoogleMapImporter

        return GoogleMapImporter()

    if str(qr_type_name) == "e050984a-4b7d-49d2-b92c-ee2d0fc7ceb9":
        from .importers.whatsapp import WhatsappImporter

        return WhatsappImporter()

    if str(qr_type_name) == "c5cf8fdb-6a63-4143-9a91-53b15b70a97e":
        from .importers.website import WebsiteImporter

        return WebsiteImporter()

    raise ValueError(
        f"Bulk import is not supported for QR type: {qr_type.name}"
    )


def get_bulk_upload_limit(user):
    subscription = Subscription.get_usage_subscription_for_user(user)
    if not subscription:
        return 0

    if subscription.bulk_upload_limit is not None:
        return subscription.bulk_upload_limit

    package_plan = getattr(subscription, "package_plan", None)
    return package_plan.max_bulk_upload if package_plan else 0


def get_qr_limit(user):
    subscription = Subscription.get_usage_subscription_for_user(user)
    if not subscription:
        return 0

    if subscription.qr_limit is not None:
        return subscription.qr_limit

    package_plan = getattr(subscription, "package_plan", None)
    return package_plan.max_qrs if package_plan else 0


def load_import_workbook_rows(import_job, max_rows=None):
    file_obj = getattr(import_job, "file", None)
    if not file_obj:
        raise ValueError("Import file is missing.")

    # Read from the underlying file object instead of relying on .path.
    # This works for both unsaved uploaded files and persisted storage-backed files.
    source = getattr(file_obj, "file", file_obj)
    if hasattr(source, "seek"):
        source.seek(0)

    workbook = load_workbook(
        source,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError("Excel file is empty.")

    headers = [
        str(header).strip()
        if header is not None
        else ""
        for header in rows[0]
    ]
    data_rows = rows[1:]

    if max_rows is not None and len(data_rows) > max_rows:
        raise ValueError(f"Excel file cannot contain more than {max_rows} rows for your package.")

    return headers, data_rows

@shared_task
def process_qr_import(import_job_id):
    try:
        import_job = QRImportJob.objects.get(
            id=import_job_id
        )

        # TODO:
        # Get PROCESSING status from ConfigChoice
        # import_job.status = processing_status

        import_job.started_at = timezone.now()
        import_job.save(
            update_fields=["started_at", "updated_at"]
        )

        bulk_upload_limit = get_bulk_upload_limit(import_job.user)
        headers, data_rows = load_import_workbook_rows(
            import_job,
            max_rows=bulk_upload_limit,
        )

        qr_limit = get_qr_limit(import_job.user)
        if qr_limit is not None:
            existing_qr_count = QRCode.objects.filter(
                created_by=import_job.user,
                is_deleted=False,
            ).count()
            if existing_qr_count + len(data_rows) > qr_limit:
                raise ValueError(
                    f"QR code limit reached for your package. "
                    f"Remaining QR slots: {max(qr_limit - existing_qr_count, 0)}."
                )

        import_job.total_rows = len(data_rows)
        import_job.save(
            update_fields=[
                "total_rows",
                "updated_at",
            ]
        )

        # Get importer based on QR type
        print("🔥 Getting importer")
        importer = get_importer(import_job.qr_type)

        print("🔥 Importer obtained:", importer)

        # Validate Excel headers
        print("🔥 Headers:", headers)
        print("🔥 Starting header validation")

        importer.validate_headers(headers)

        print("🔥 Header validation completed")

        # Process rows
        print("🔥 Number of data rows:", len(data_rows))

        # Process rows
        for index, row_values in enumerate(
                data_rows,
                start=2,
        ):
            try:
                row = dict(
                    zip(headers, row_values)
                )

                print(f"🔥 Processing row {index}: {row}")

                with transaction.atomic():
                    importer.process_row(
                        row=row,
                        job=import_job,
                        row_number=index,
                    )

                import_job.successful_rows += 1

                print(f"✅ Row {index} processed successfully")

            except Exception as exc:
                print(
                    f"❌ Row {index} failed: {repr(exc)}"
                )

                import_job.failed_rows += 1

                if not import_job.error_details:
                    import_job.error_details = []

                import_job.error_details.append(
                    {
                        "row": index,
                        "error": str(exc),
                    }
                )

            import_job.processed_rows += 1

            if import_job.total_rows > 0:
                import_job.progress = int(
                    (
                            import_job.processed_rows
                            / import_job.total_rows
                    ) * 100
                )
            else:
                import_job.progress = 100

            import_job.save(
                update_fields=[
                    "processed_rows",
                    "successful_rows",
                    "failed_rows",
                    "progress",
                    "error_details",
                    "updated_at",
                ]
            )

        # Import completed
        import_job.completed_at = timezone.now()
        import_job.progress = 100

        import_job.save(
            update_fields=[
                "completed_at",
                "progress",
                "updated_at",
            ]
        )

    except QRImportJob.DoesNotExist:
        return

    except ValueError as exc:
        QRImportJob.objects.filter(
            id=import_job_id
        ).update(
            error_message=str(exc),
            completed_at=timezone.now(),
        )
        return

    except Exception as exc:
        print(
            f"❌ QR import job {import_job_id} failed: "
            f"{repr(exc)}"
        )

        QRImportJob.objects.filter(
            id=import_job_id
        ).update(
            error_message=str(exc),
            completed_at=timezone.now(),
        )

        raise


# tasks.py
from celery import shared_task
from .services.domain_verification import DomainVerificationService
from .models import CustomDomain
import logging

logger = logging.getLogger(__name__)


@shared_task(bind=True, max_retries=3)
def verify_and_activate_domain_async(self, domain_id):
    """Asynchronously verify and activate a domain"""
    try:
        domain = CustomDomain.objects.get(id=domain_id, is_deleted=False)

        verification_service = DomainVerificationService()
        result = verification_service.verify_and_activate_domain(domain)

        if not result['success']:
            # Retry with exponential backoff
            raise Exception(f"Domain verification failed: {result.get('error', 'Unknown error')}")

        return result

    except CustomDomain.DoesNotExist:
        logger.error(f"Domain {domain_id} not found")
        return {'success': False, 'error': 'Domain not found'}
    except Exception as e:
        logger.error(f"Domain verification error: {str(e)}")
        # Retry with exponential backoff
        raise self.retry(exc=e, countdown=60 * (2 ** self.request.retries))


@shared_task
def renew_expiring_ssl_certificates():
    """Renew SSL certificates that are about to expire"""
    from datetime import timedelta
    from django.utils import timezone
    from .services.ssl_provisioning_service import SSLProvisioningService

    # Find domains with SSL expiring in less than 30 days
    expiry_threshold = timezone.now() + timedelta(days=30)
    domains = CustomDomain.objects.filter(
        status=CustomDomain.Status.ACTIVE,
        ssl_expires_at__lte=expiry_threshold,
        is_deleted=False
    )

    results = []
    for domain in domains:
        ssl_service = SSLProvisioningService(domain.domain)
        result = ssl_service.renew_certificate()

        if result['success']:
            # Update SSL expiry
            from datetime import datetime
            import subprocess
            cmd = ['openssl', 'x509', '-in', f'/etc/letsencrypt/live/{domain.domain}/fullchain.pem', '-enddate',
                   '-noout']
            proc = subprocess.run(cmd, capture_output=True, text=True)
            if proc.returncode == 0:
                expiry_str = proc.stdout.strip().replace('notAfter=', '')
                domain.ssl_expires_at = datetime.strptime(expiry_str, '%b %d %H:%M:%S %Y %Z')
                domain.save()

        results.append({
            'domain': domain.domain,
            'success': result['success'],
            'error': result.get('error') if not result['success'] else None
        })

    return {
        'processed': len(domains),
        'results': results
    }
